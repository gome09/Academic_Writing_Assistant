# -*- coding: utf-8 -*-
"""
多格式读取器 —— 读取 docx / pdf / txt / md / json，
输出统一的中间结构 Document。

契约由下方 Block / Document TypedDict 显式定义（T4-5），运行时仍为普通
dict，TypedDict 仅作类型层契约与文档，完全向后兼容。历史隐式 dict 调用
（b["kind"] / d["blocks"] 等）无需改动。

设计原则：
  - 任何一种格式失败不影响其它格式；缺库时给出清晰提示。
  - 尽量从原文档识别标题层级（docx 的 Heading 样式 / md 的 #）。
"""
from __future__ import annotations
import json
import os
import re
from typing import Any, Callable, NotRequired, TypedDict


# ---------------------------------------------------------------------------
#  统一数据契约（T4-5）
#  运行时为普通 dict；此处 TypedDict 仅声明字段契约，便于类型检查与文档化。
#  iter_node_blocks（organizer.py）兼容层仍按 dict 消费，可逐步迁移。
# ---------------------------------------------------------------------------
class Block(TypedDict):
    """内容块契约。kind/level/text 为必填；rows/data/ext 按需出现。"""

    kind: str           # heading|paragraph|list_item|table|code|image
    level: int          # heading 级别，非标题为 0
    text: str           # 纯文本（已 strip）；table/image 可为空串
    # 以下字段 NotRequired：仅特定 kind 出现，运行时可能缺省
    rows: NotRequired[Any]            # 仅 table：[[cell, ...], ...]
    data: NotRequired[Any]            # 仅 image：原始字节
    ext: NotRequired[Any]             # 仅 image：扩展名（如 ".png"）


class Document(TypedDict):
    """读取器统一输出。"""

    source: str         # 文件路径
    type: str           # docx|pdf|txt|md|json|xlsx|csv|image
    blocks: list[Block]  # 顺序保留
    meta: dict[str, Any]  # 可选：从 json/frontmatter/core_properties 提取的元信息


# ---------------------------------------------------------------------------
#  工具
# ---------------------------------------------------------------------------
def _block(kind, text="", level=0, rows=None) -> Block:
    return {"kind": kind, "level": level, "text": text.strip(), "rows": rows}


def _clean(s: str) -> str:
    return re.sub(r"[ \t]+", " ", (s or "")).strip()


def _read_text(path: str) -> str:
    """读取文本文件。

    三步回退顺序固化（不要交换）：
      1) utf-8-sig
      2) gb18030
      3) utf-8 (errors="replace")

    为什么：UTF-8 文本几乎不会被 gb18030 抢先命中（utf-8 先试且
    严格）；GBK/GB18030 文本解码 utf-8 必报 UnicodeDecodeError 落到第二档。
    打乱码文件由第三步 errors="replace" 兜底，不会抛异常。
    """
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# ---------------------------------------------------------------------------
#  读取器注册（T4-3 插件化）
#  _READERS 为扩展名 -> 读取函数 的注册表；第三方可用 register_reader 装饰器
#  注册新读取器，无需修改核心分发逻辑。向后兼容：_READERS 仍是普通 dict，
#  read_dir_detailed 的扩展名过滤与历史行为一致。
# ---------------------------------------------------------------------------
_READERS: dict[str, Callable[..., Document]] = {}


def register_reader(*exts: str):
    """装饰器：将读取函数注册到指定扩展名（T4-3）。

    用法：
        @register_reader(".md", ".markdown")
        def read_md(path: str) -> Document: ...
    注册后 read_file / read_dir_detailed 自动识别该扩展名。
    """

    def decorator(fn):
        for ext in exts:
            _READERS[ext] = fn
        return fn

    return decorator


# ---------------------------------------------------------------------------
#  TXT
# ---------------------------------------------------------------------------
@register_reader(".txt", ".text")
def read_txt(path: str) -> Document:
    raw = _read_text(path)
    blocks = []
    for para in re.split(r"\n\s*\n", raw):
        para = para.strip()
        if para:
            blocks.append(_block("paragraph", para))
    return {"source": path, "type": "txt", "blocks": blocks, "meta": {}}


# ---------------------------------------------------------------------------
#  Markdown
# ---------------------------------------------------------------------------
@register_reader(".md", ".markdown")
def read_md(path: str) -> Document:
    raw = _read_text(path)

    meta = {}
    # YAML frontmatter
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", raw, re.DOTALL)
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip().strip("\"'")
        raw = raw[m.end():]

    blocks = []
    lines = raw.splitlines()
    i = 0
    para_buf = []

    def flush_para():
        if para_buf:
            blocks.append(_block("paragraph", " ".join(para_buf)))
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        # ATX 标题
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h:
            flush_para()
            blocks.append(_block("heading", h.group(2), level=len(h.group(1))))
            i += 1
            continue
        # 列表项
        li = re.match(r"^\s*(?:[-*+]|\d+\.)\s+(.*)$", line)
        if li:
            flush_para()
            blocks.append(_block("list_item", li.group(1)))
            i += 1
            continue
        # 代码块
        if line.strip().startswith("```"):
            flush_para()
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            blocks.append(_block("code", "\n".join(buf)))
            i += 1
            continue
        # 空行 -> 段落边界
        if not line.strip():
            flush_para()
            i += 1
            continue
        para_buf.append(line.strip())
        i += 1
    flush_para()
    return {"source": path, "type": "md", "blocks": blocks, "meta": meta}


# ---------------------------------------------------------------------------
#  JSON —— 支持两类：结构化大纲 或 任意数据
# ---------------------------------------------------------------------------
@register_reader(".json")
def read_json(path: str) -> Document:
    data = json.loads(_read_text(path))

    blocks = []
    meta = {}

    def walk(obj, level=1):
        if isinstance(obj, dict):
            # 约定字段：title/heading + content/text + children/sections
            title = obj.get("title") or obj.get("heading")
            if title:
                blocks.append(_block("heading", str(title), level=min(level, 3)))
            body = obj.get("content") or obj.get("text") or obj.get("body")
            if body:
                if isinstance(body, list):
                    for it in body:
                        blocks.append(_block("paragraph", str(it)))
                else:
                    blocks.append(_block("paragraph", str(body)))
            for kids_key in ("children", "sections", "subsections", "items"):
                if isinstance(obj.get(kids_key), list):
                    for kid in obj[kids_key]:
                        walk(kid, level + 1)
            # 其它标量字段收进 meta（仅顶层）
            if level == 1:
                for k, v in obj.items():
                    if k in ("title", "heading", "content", "text", "body",
                             "children", "sections", "subsections", "items"):
                        continue
                    if isinstance(v, (str, int, float)):
                        meta[k] = v
        elif isinstance(obj, list):
            for it in obj:
                walk(it, level)
        else:
            blocks.append(_block("paragraph", str(obj)))

    walk(data, 1)
    return {"source": path, "type": "json", "blocks": blocks, "meta": meta}


# ---------------------------------------------------------------------------
#  DOCX
# ---------------------------------------------------------------------------
@register_reader(".docx")
def read_docx(path: str) -> Document:
    try:
        import docx  # python-docx
    except ImportError:
        raise RuntimeError("需要 python-docx：pip install python-docx")

    doc = docx.Document(path)
    blocks = []
    # 按 body 的 XML 子元素顺序混排段落与表格，保留表格在原文中的位置
    # （body 里还有 w:sectPr 等其它元素，跳过即可）。
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for el in doc.element.body:
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            # 段内图片：a:blip 的 r:embed 指向图片关系部件。
            # 必须先提图片再做空文本跳过——图片段落通常没有文字。
            for blip in el.findall(".//" + qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                part = doc.part.related_parts.get(rid) if rid else None
                if part is None:
                    continue
                b = _block("image")
                b["data"] = part.blob
                b["ext"] = os.path.splitext(str(part.partname))[1].lower() or ".png"
                blocks.append(b)
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name or "").lower()
            m = re.search(r"heading\s*(\d)", style)
            if m:
                blocks.append(_block("heading", text, level=int(m.group(1))))
            elif style.startswith("list") or style.startswith("bullet"):
                blocks.append(_block("list_item", text))
            elif _looks_like_manual_heading(text):
                hm = _PDF_HEADING.match(text)
                blocks.append(_block("heading", text,
                                     level=_pdf_heading_level(hm.group(1))))
            else:
                blocks.append(_block("paragraph", text))
        elif el.tag == qn("w:tbl"):
            t = Table(el, doc)
            rows = []
            for row in t.rows:
                rows.append([_clean(c.text) for c in row.cells])
            if rows:
                blocks.append(_block("table", "", rows=rows))

    meta = {}
    cp = doc.core_properties
    if cp.title:
        meta["title"] = cp.title
    if cp.author:
        meta["author"] = cp.author
    return {"source": path, "type": "docx", "blocks": blocks, "meta": meta}


# ---------------------------------------------------------------------------
#  PDF
# ---------------------------------------------------------------------------
def _join_lines(lines):
    """合并断行：ASCII 单词之间补空格，中文直接相连。"""
    out = ""
    for ln in lines:
        if (out and out[-1].isascii() and out[-1].isalnum()
                and ln[:1].isascii() and ln[:1].isalnum()):
            out += " "
        out += ln
    return out


def _pdf_heading_level(prefix: str) -> int:
    if "章" in prefix:
        return 1
    m = re.match(r"\d+(\.\d+)*", prefix)
    return min(m.group(0).count(".") + 1, 3) if m else 1


_PDF_HEADING = re.compile(
    r"^(第\s*[一二三四五六七八九十百\d]+\s*章"
    r"|(?!\d{4}(?:[\s年]|$))\d+(\.\d+)*[\s、.．])")


def _looks_like_manual_heading(text: str) -> bool:
    """无标题样式但形如标题：编号开头 + 短（≤25字）+ 不含句中/句末标点。"""
    if len(text) > 25 or re.search(r"[。！？；，,;]", text):
        return False
    m = _PDF_HEADING.match(text)
    if not m:
        return False
    # 「第X章」后须跟空白或行尾，避免「第一章正文」这类连写句子被误判为标题
    # （数字编号分支的正则本身已要求编号后带分隔符，无需再查）。
    if "章" in m.group(1):
        rest = text[m.end():]
        return not rest or rest[0].isspace()
    return True


def _pdf_lines_to_blocks(txt: str, blocks: list) -> None:
    """PDF 文本按行重组：行尾终止标点断段；短编号行识别为标题。

    PDF 提取的文本几乎没有连续空行，不能按空行分段；
    这里以「句末标点在行尾」为段落边界，是对中文论文的合理近似。
    """
    buf = []

    def flush():
        if buf:
            blocks.append(_block("paragraph", _join_lines(buf)))
            buf.clear()

    for line in txt.splitlines():
        line = _clean(line)
        if not line:
            flush()
            continue
        m = _PDF_HEADING.match(line)
        if (m and len(line) <= 25
                and not re.search(r"[。！？；，,;:.]\s*$", line)):
            flush()
            blocks.append(_block("heading", line,
                                 level=_pdf_heading_level(m.group(1))))
            continue
        buf.append(line)
        if re.search(r"[。！？!?]$", line):
            flush()
    flush()


def _ocr_available() -> bool:
    """检查 OCR 依赖（pytesseract + pdf2image）是否可用。"""
    try:
        import pytesseract  # noqa: F401
        import pdf2image  # noqa: F401
        return True
    except ImportError:
        return False


def _ocr_pdf(path: str) -> list:
    """用 Tesseract OCR 识别扫描件 PDF，返回文本块列表。

    依赖 pytesseract + pdf2image（可选依赖，未安装时抛 RuntimeError）。
    """
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError as exc:
        raise RuntimeError(
            "OCR 需要安装 pytesseract 和 pdf2image："
            "pip install pytesseract pdf2image（还需系统安装 Tesseract-OCR）"
        ) from exc
    blocks = []
    images = convert_from_path(path)
    for img in images:
        txt = pytesseract.image_to_string(img, lang="chi_sim+eng")
        if txt.strip():
            _pdf_lines_to_blocks(txt, blocks)
    return blocks


def _ensure_has_text(blocks: list, path: str, ocr: bool = False) -> None:
    """扫描件（图片型 PDF）提取不到文字时处理。

    ocr=True 时尝试 OCR 识别；否则或 OCR 失败时报错。
    """
    if any(b.get("text") for b in blocks):
        return
    if any(b.get("kind") == "table" for b in blocks):
        return
    if any(b.get("kind") == "image" for b in blocks):
        return  # T4-2：已提取内嵌图片，视为有内容
    if ocr:
        try:
            ocr_blocks = _ocr_pdf(path)
            if any(b.get("text") for b in ocr_blocks):
                blocks.extend(ocr_blocks)
                print(f"  [OCR] {os.path.basename(path)}：OCR 识别成功，"
                      f"提取到 {len(ocr_blocks)} 个文本块。")
                return
            print(f"  [OCR] {os.path.basename(path)}：OCR 未识别到文字。")
        except RuntimeError as e:
            print(f"  [OCR] {os.path.basename(path)}：{e}")
        except Exception as e:  # noqa: BLE001
            print(f"  [OCR] {os.path.basename(path)}：OCR 失败：{e}")
    raise RuntimeError(
        "未提取到任何文字，可能是扫描件（图片型 PDF）。"
        "请先用 OCR 工具（如 WPS/Acrobat/umi-ocr）转成可复制文本的 PDF 再试，"
        "或加 --ocr 启用内置 OCR（需安装 pytesseract + pdf2image）")


def _extract_pdf_images(path: str) -> tuple[list, int]:
    """提取 PDF 内嵌图片为 image 块列表（T4-2）。

    依赖 pypdf（可选依赖，未安装时返回 ([], -1) 标识不可用）。
    pypdf 的 page.images 访问器返回 ImageFile，含 .data（原始字节）与 .name。
    单张图提取失败不影响其它图。
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        return [], -1
    blocks = []
    count = 0
    try:
        reader = PdfReader(path)
        for page in reader.pages:
            try:
                images = page.images
            except Exception:  # noqa: BLE001 —— 部分页无图片资源对象
                continue
            for img in images:
                try:
                    name = getattr(img, "name", "") or ""
                    ext = os.path.splitext(name)[1].lower() or ".png"
                    b = _block("image")
                    b["data"] = img.data
                    b["ext"] = ext
                    blocks.append(b)
                    count += 1
                except Exception:  # noqa: BLE001 —— 跳过无法解码的单张图
                    continue
    except Exception:  # noqa: BLE001 —— 整体失败返回已提取部分
        pass
    return blocks, count


@register_reader(".pdf")
def read_pdf(path: str, ocr: bool = False, extract_images: bool = False) -> Document:
    blocks = []
    meta = {}
    try:
        import pdfplumber
    except ImportError:
        pdfplumber = None

    img_count = 0
    if pdfplumber is not None:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                img_count += len(page.images)
                # 表格（find_tables 以便拿到 bbox 用于正文去重）
                tables = page.find_tables()
                bboxes = []
                for tbl in tables:
                    rows = [[(_clean(c) if c else "") for c in row]
                            for row in (tbl.extract() or [])]
                    rows = [r for r in rows if any(r)]
                    if rows:
                        blocks.append(_block("table", "", rows=rows))
                        bboxes.append(tbl.bbox)

                # 正文：过滤掉落在表格 bbox 内的字符，避免内容重复
                def _outside(obj, _bx=tuple(bboxes)):
                    cx = (obj["x0"] + obj["x1"]) / 2
                    cy = (obj["top"] + obj["bottom"]) / 2
                    return not any(x0 <= cx <= x1 and y0 <= cy <= y1
                                   for (x0, y0, x1, y1) in _bx)

                target = page.filter(_outside) if bboxes else page
                txt = target.extract_text() or ""
                _pdf_lines_to_blocks(txt, blocks)
    else:
        # 退化到 pypdf
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("需要 pdfplumber 或 pypdf：pip install pdfplumber")
        reader = PdfReader(path)
        for page in reader.pages:
            txt = page.extract_text() or ""
            _pdf_lines_to_blocks(txt, blocks)

    # T4-2：可选提取 PDF 内嵌图片为 image 块（默认关，因可能量大）
    if extract_images:
        img_blocks, got = _extract_pdf_images(path)
        if got < 0:
            print(f"  [提示] {os.path.basename(path)}：未安装 pypdf，"
                  "无法提取内嵌图片（pip install pypdf）")
        elif got:
            blocks.extend(img_blocks)
            print(f"  [图片] {os.path.basename(path)}：提取 {got} 张内嵌图片为 image 块")
        elif img_count:
            print(f"  [提示] {os.path.basename(path)}：检测到 {img_count} 张图片，"
                  "但未能提取为 image 块（可能为矢量或掩码图）")
    elif img_count:
        print(f"  [提示] {os.path.basename(path)}：检测到 {img_count} 张图片，"
              "PDF 图片暂不导入；加 --extract-pdf-images 可提取内嵌图片")

    _ensure_has_text(blocks, path, ocr=ocr)
    return {"source": path, "type": "pdf", "blocks": blocks, "meta": meta}


# ---------------------------------------------------------------------------
#  XLSX / CSV
# ---------------------------------------------------------------------------
@register_reader(".xlsx")
def read_xlsx(path: str) -> Document:
    """每个非空工作表 -> 一个 table 块。合并单元格的非锚点格读出 None -> ""。

    data_only=True 读公式的缓存计算值；文件从未被 Excel 打开过时可能为
    None，同样落为空串，属可接受降级。
    """
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("需要 openpyxl：pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    blocks = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else _clean(str(c)) for c in row]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append(_block("table", "", rows=rows))
    finally:
        wb.close()
    if not blocks:
        raise RuntimeError("工作簿中没有任何非空工作表")
    return {"source": path, "type": "xlsx", "blocks": blocks, "meta": {}}


@register_reader(".csv")
def read_csv(path: str) -> Document:
    """整个 CSV -> 单个 table 块；全空行跳过。编码回退复用 _read_text。"""
    import csv
    import io
    raw = _read_text(path)
    rows = []
    for row in csv.reader(io.StringIO(raw)):
        cells = [_clean(c) for c in row]
        if any(cells):
            rows.append(cells)
    if not rows:
        raise RuntimeError("CSV 中没有任何非空行")
    return {"source": path, "type": "csv",
            "blocks": [_block("table", "", rows=rows)], "meta": {}}


# ---------------------------------------------------------------------------
#  独立图片（截图等）
# ---------------------------------------------------------------------------
@register_reader(".png", ".jpg", ".jpeg", ".bmp", ".webp")
def read_image(path: str) -> Document:
    """整个文件 -> 单个 image 块（原始字节 + 扩展名），供插图与视觉理解。"""
    with open(path, "rb") as f:
        data = f.read()
    if not data:
        raise RuntimeError("图片文件为空")
    b = _block("image")
    b["data"] = data
    b["ext"] = os.path.splitext(path)[1].lower() or ".png"
    return {"source": path, "type": "image", "blocks": [b], "meta": {}}


# ---------------------------------------------------------------------------
#  分发（T4-3：读取器在定义处用 @register_reader 注册，此处无需静态字典）
# ---------------------------------------------------------------------------


def read_file(path: str, ocr: bool = False, extract_images: bool = False) -> Document:
    """按扩展名读取单个文件，返回统一 Document。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return read_pdf(path, ocr=ocr, extract_images=extract_images)
    reader = _READERS.get(ext)
    if reader is None:
        raise ValueError(f"不支持的文件类型：{ext}（{path}）")
    return reader(path)


def read_dir_detailed(dir_path: str, ocr: bool = False,
                      extract_images: bool = False):
    """Return (documents, errors) while keeping per-file failure details."""
    docs = []
    errors = []
    for name in sorted(os.listdir(dir_path)):
        full = os.path.join(dir_path, name)
        if not os.path.isfile(full):
            continue
        if os.path.splitext(name)[1].lower() not in _READERS:
            continue
        try:
            docs.append(read_file(full, ocr=ocr, extract_images=extract_images))
            print(f"  [读取] {name}")
        except Exception as e:  # noqa: BLE001
            print(f"  [跳过] {name}: {e}")
            errors.append((full, str(e)))
    return docs, errors
